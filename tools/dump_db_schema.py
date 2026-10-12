#!/usr/bin/env python3
"""
Dump every table/column in the configured JobBOSS database to an Excel file.

Read-only -- only ever queries INFORMATION_SCHEMA and sys.* system catalog
views, never application data. Built for schema discovery (e.g. finding the
real table behind a JobBOSS canned report) when there's no Crystal Reports
definition or SSMS access handy.

Credentials: by default, reuses whatever's already configured in JobDocs
itself (Settings > Advanced Settings -- reads settings.json + the OS
credential store, nothing to re-enter). Override with JOBBOSS_DB_* env vars
if you want to run this with different/broader credentials just for this
dump -- useful because the app's own read-only login is deliberately scoped
to only a few tables (Job/Delivery/Job_Operation), and SQL Server hides
INFORMATION_SCHEMA/sys.tables rows for anything that login has no permission
on. Two ways to see everything:
  1. Run this with a login that already has broader access (e.g. your own
     SSMS/DBA login): JOBBOSS_DB_USER=... JOBBOSS_DB_PASS=... python tools/dump_db_schema.py
  2. Or grant metadata-only visibility to the existing read-only login
     (does NOT expose any row data, just table/column names and types):
         GRANT VIEW DEFINITION ON DATABASE::<your_jobboss_database> TO jobdocs_reporting_ro;

Usage:
    python tools/dump_db_schema.py [output.xlsx]
"""
import json
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytds  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

from experimental.db_integration import _connect_kwargs, is_configured  # noqa: E402
from shared.utils import get_config_dir  # noqa: E402

_TABLES_QUERY = """
SELECT TABLE_SCHEMA, TABLE_NAME, TABLE_TYPE
FROM INFORMATION_SCHEMA.TABLES
ORDER BY TABLE_SCHEMA, TABLE_NAME
"""

_COLUMNS_QUERY = """
SELECT TABLE_SCHEMA, TABLE_NAME, COLUMN_NAME, ORDINAL_POSITION, DATA_TYPE,
       CHARACTER_MAXIMUM_LENGTH, IS_NULLABLE
FROM INFORMATION_SCHEMA.COLUMNS
ORDER BY TABLE_SCHEMA, TABLE_NAME, ORDINAL_POSITION
"""

# sys.partitions, not COUNT(*) per table -- fast (reads stored row-count
# stats, no table scans) even against a large JobBOSS database. Useful for
# spotting which tables are actually populated (same signal that confirmed
# SO_Detail was unused -- 0 rows -- earlier in this project, back then via a
# real COUNT(*) against one already-known table; this generalizes that check
# to every table at once for a first-pass discovery sweep).
_ROW_COUNTS_QUERY = """
SELECT s.name AS TABLE_SCHEMA, t.name AS TABLE_NAME, SUM(p.rows) AS APPROX_ROWS
FROM sys.tables t
JOIN sys.schemas s ON s.schema_id = t.schema_id
JOIN sys.partitions p ON p.object_id = t.object_id AND p.index_id IN (0, 1)
GROUP BY s.name, t.name
"""


def _load_settings() -> dict:
    """JOBBOSS_DB_* env vars win if set; otherwise fall back to JobDocs' own
    stored settings.json + keyring password, same as the running app uses.
    """
    env_host = os.environ.get('JOBBOSS_DB_HOST', '').strip()
    if env_host:
        return {
            'jobboss_db_host': env_host,
            'jobboss_db_port': os.environ.get('JOBBOSS_DB_PORT', '').strip(),
            'jobboss_db_name': os.environ.get('JOBBOSS_DB_NAME', '').strip(),
            'jobboss_db_user': os.environ.get('JOBBOSS_DB_USER', '').strip(),
        }
    settings_file = get_config_dir() / 'settings.json'
    if not settings_file.exists():
        raise SystemExit(
            f"No settings.json at {settings_file} and no JOBBOSS_DB_HOST env var set.\n"
            f"Either run JobDocs and configure the JobBOSS DB connection in "
            f"Settings > Advanced Settings first, or set JOBBOSS_DB_HOST/PORT/NAME/USER "
            f"(+ JOBBOSS_DB_PASS, since the env-var path doesn't go through keyring)."
        )
    return json.loads(settings_file.read_text(encoding='utf-8'))


def _get_password(settings: dict) -> str:
    env_pass = os.environ.get('JOBBOSS_DB_PASS', '')
    if env_pass:
        return env_pass
    from experimental.db_integration import get_password
    return get_password(settings.get('jobboss_db_user', ''))


def _autofit(ws):
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col_cells in ws.columns:
        width = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
        ws.column_dimensions[col_cells[0].column_letter].width = max(10, min(40, width + 2))


def main():
    settings = _load_settings()
    if not is_configured(settings) and not os.environ.get('JOBBOSS_DB_PASS'):
        raise SystemExit(
            "JobBOSS DB connection isn't fully configured (missing host/name/user/password). "
            "Configure it in JobDocs' Settings, or set JOBBOSS_DB_HOST/PORT/NAME/USER/PASS."
        )

    out_path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path('jobboss_schema_dump.xlsx')

    dsn, connect_kwargs = _connect_kwargs(settings)
    connect_kwargs['password'] = _get_password(settings)
    print(f"Connecting to {settings.get('jobboss_db_host')} / {settings.get('jobboss_db_name')} ...")

    with pytds.connect(dsn, **connect_kwargs) as conn:
        with conn.cursor() as cur:
            cur.execute(_TABLES_QUERY)
            tables = cur.fetchall()
            cur.execute(_COLUMNS_QUERY)
            columns = cur.fetchall()
            cur.execute(_ROW_COUNTS_QUERY)
            row_counts = {(r['TABLE_SCHEMA'], r['TABLE_NAME']): r['APPROX_ROWS'] for r in cur.fetchall()}

    print(f"Found {len(tables)} tables, {len(columns)} columns.")
    if not tables:
        print(
            "Zero tables visible -- this login likely lacks metadata visibility "
            "for anything it hasn't been granted SELECT on. See the module "
            "docstring for how to fix that (GRANT VIEW DEFINITION, or use a "
            "broader login via JOBBOSS_DB_* env vars)."
        )

    wb = Workbook()
    ws_tables = wb.active
    ws_tables.title = "Tables"
    ws_tables.append(["Schema", "Table", "Type", "Approx Rows"])
    for row in tables:
        key = (row['TABLE_SCHEMA'], row['TABLE_NAME'])
        ws_tables.append([row['TABLE_SCHEMA'], row['TABLE_NAME'], row['TABLE_TYPE'], row_counts.get(key)])

    ws_cols = wb.create_sheet("Columns")
    ws_cols.append(["Schema", "Table", "Column", "Position", "Data Type", "Max Length", "Nullable"])
    for row in columns:
        ws_cols.append([
            row['TABLE_SCHEMA'], row['TABLE_NAME'], row['COLUMN_NAME'], row['ORDINAL_POSITION'],
            row['DATA_TYPE'], row['CHARACTER_MAXIMUM_LENGTH'], row['IS_NULLABLE'],
        ])

    for ws in (ws_tables, ws_cols):
        _autofit(ws)

    wb.save(out_path)
    print(f"Wrote {out_path.resolve()}")


if __name__ == '__main__':
    main()
