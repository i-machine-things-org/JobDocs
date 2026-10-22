"""Tests for the "By Department" export-split feature in
modules/reporting/module.py: _group_rows_by_department(), _safe_sheet_name(),
_department_average_row(), and _write_xlsx_report_sheet().

All four are pure functions (or, for _write_xlsx_report_sheet, operate on a
plain openpyxl worksheet with no PyQt involved) -- no QApplication needed,
same reasoning as _neutralize_formula_cell()/_numeric_or_text() in
test_reporting_module_export_safety.py.

Context: a first attempt at "filter Employee Efficiency by department"
(an IN-clause query filter + multi-select picker) worked in isolation but
the user wanted something simpler and more useful -- a single "By
Department" checkbox that splits the xlsx export into one worksheet tab
per department, each with a trailing Department Average row, instead of
restricting which rows come back at all. See modules/reporting/reports/
employee_efficiency.py's docstring and base.ReportDefinition.
uses_department_tabs.
"""

from openpyxl import Workbook

from modules.reporting.module import (
    _department_average_row, _group_rows_by_department, _safe_sheet_name, _write_xlsx_report_sheet,
)
from modules.reporting.reports.base import ReportDefinition


def test_group_rows_by_department_preserves_first_seen_order():
    rows = [
        ['Adams', 'MACHINING'],
        ['Baker', 'WELDING'],
        ['Carter', 'MACHINING'],
    ]
    groups = _group_rows_by_department(rows, department_col=1)
    assert [dept for dept, _ in groups] == ['MACHINING', 'WELDING']
    assert groups[0][1] == [['Adams', 'MACHINING'], ['Carter', 'MACHINING']]
    assert groups[1][1] == [['Baker', 'WELDING']]


def test_group_rows_by_department_gives_blank_department_its_own_group():
    rows = [['Adams', ''], ['Baker', 'WELDING']]
    groups = _group_rows_by_department(rows, department_col=1)
    assert groups[0] == ('(No Department)', [['Adams', '']])
    assert groups[1] == ('WELDING', [['Baker', 'WELDING']])


def test_safe_sheet_name_strips_invalid_excel_characters():
    used = set()
    assert _safe_sheet_name('CNC/LATHE', used) == 'CNC_LATHE'


def test_safe_sheet_name_truncates_to_31_chars():
    used = set()
    name = _safe_sheet_name('A' * 40, used)
    assert len(name) == 31
    assert name == 'A' * 31


def test_safe_sheet_name_dedupes_collisions():
    used = set()
    first = _safe_sheet_name('MACHINING', used)
    second = _safe_sheet_name('MACHINING', used)
    assert first == 'MACHINING'
    assert second == 'MACHINING_2'
    assert first != second


def test_safe_sheet_name_falls_back_when_name_is_blank_after_cleaning():
    used = set()
    assert _safe_sheet_name('///', used) == '___'
    assert _safe_sheet_name('', set()) == 'Sheet'


def test_department_average_row_labels_first_column_and_averages_numeric_columns():
    # Employee/text column (0) is never averaged; two numeric columns (1, 2)
    # average cleanly; mixed column (3) has one non-numeric entry and must
    # stay blank rather than averaging only the numeric subset.
    rows = [
        ['Adams', '10', '75.0', 'N/A'],
        ['Baker', '20', '85.0', '5'],
    ]
    avg = _department_average_row(rows)
    assert avg[0] == 'Department Average'
    assert avg[1] == 15.0
    assert avg[2] == 80.0
    assert avg[3] == ''


def test_department_average_row_empty_input_returns_empty():
    assert _department_average_row([]) == []


def test_department_average_row_excludes_blank_cells_without_voiding_the_column():
    # A blank cell ('') means "not scoreable" (e.g. an Indirect operation's
    # blank Adj Hrs) -- it must be excluded from the average, not void the
    # whole column the way a genuine non-numeric anomaly ('N/A') still does.
    rows = [
        ['Adams', '10.0', ''],
        ['Baker', '20.0', '5.0'],
    ]
    avg = _department_average_row(rows)
    assert avg[1] == 15.0
    assert avg[2] == 5.0  # averaged over just the one real value, not voided


def test_department_average_row_restricts_to_average_columns():
    # Column 1 ('job') is all-numeric by coincidence but excluded from
    # average_columns -- must stay blank even though it would otherwise
    # pass the numeric check. Column 2 ('hours') is included and averages.
    rows = [
        ['Adams', '30101', '10.0'],
        ['Baker', '30102', '20.0'],
    ]
    avg = _department_average_row(rows, average_columns={2})
    assert avg[1] == ''
    assert avg[2] == 15.0


def _report_def(**overrides):
    defaults = dict(
        name="Test Report", columns=('employee', 'department', 'hours'),
        headers=['Employee', 'Department', 'Hours'], fetch=lambda settings, filters: [],
    )
    defaults.update(overrides)
    return ReportDefinition(**defaults)


def test_write_xlsx_report_sheet_includes_average_row_and_table():
    wb = Workbook()
    ws = wb.active
    data_rows = [['Adams', 'MACHINING', '10.00'], ['Baker', 'MACHINING', '20.00']]
    _write_xlsx_report_sheet(
        ws, title_block=[['Title'], []], headers=['Employee', 'Department', 'Hours'],
        data_rows=data_rows, report_def=_report_def(),
        average_row=_department_average_row(data_rows), table_name='ReportTable',
    )
    values = [[cell.value for cell in row] for row in ws.iter_rows()]
    assert ['Department Average', '', 15.0] in values
    assert ws.tables['ReportTable'] is not None


def test_write_xlsx_report_sheet_without_average_row_matches_old_single_sheet_behavior():
    wb = Workbook()
    ws = wb.active
    data_rows = [['Adams', 'MACHINING', '10.00']]
    _write_xlsx_report_sheet(
        ws, title_block=[['Title'], []], headers=['Employee', 'Department', 'Hours'],
        data_rows=data_rows, report_def=_report_def(),
    )
    values = [[cell.value for cell in row] for row in ws.iter_rows()]
    assert all('Department Average' not in row for row in values)
