"""Tests for the "Employee Total" row feature in modules/reporting/module.py
(_employee_total_rows()), plus its interaction with _write_xlsx_report_sheet()'s
Excel Table wrapping.

Context: a user comparing our Detail export against JobBOSS's real one
pointed out that JobBOSS's native report is laid out in per-employee blocks,
each closed out by a total line and a blank spacer row before the next
employee -- ours was a flat list with no such rows at all. See
modules/reporting/reports/base.py's ReportDefinition.uses_employee_total_rows
and employee_efficiency.py's module docstring for the full writeup.
"""

from openpyxl import Workbook

from modules.reporting.module import _employee_total_rows, _write_xlsx_report_sheet
from modules.reporting.reports.base import ReportDefinition

_COLUMNS = (
    'employee', 'employee_code', 'job', 'work_center', 'department', 'operation',
    'setup_act_hrs', 'setup_est_hrs', 'setup_adj_hrs', 'setup_pct_eff',
    'run_act_hrs', 'run_est_hrs', 'run_qty', 'run_adj_hrs', 'run_pct_eff', 'pct_indirect',
)


def _report_def(**overrides):
    defaults = dict(
        name="Test Detail", columns=_COLUMNS, headers=list(_COLUMNS),
        fetch=lambda settings, filters: [], uses_employee_total_rows=True,
    )
    defaults.update(overrides)
    return ReportDefinition(**defaults)


def _row(employee, employee_code=None, setup_act='0.00', setup_est='0.00', setup_adj='', setup_pct='',
         run_act='0.00', run_est='0.00', run_qty='0', run_adj='', run_pct='', pct_indirect='5.0'):
    # employee_code defaults to the display name itself -- fine for every
    # existing test here, which already uses a unique name per employee;
    # only the duplicate-display-name test below needs to pass a real,
    # distinct code explicitly.
    code = employee_code if employee_code is not None else employee
    return [employee, code, 'J1', 'MILLS', 'DEPT', 'OP1', setup_act, setup_est, setup_adj, setup_pct,
            run_act, run_est, run_qty, run_adj, run_pct, pct_indirect]


def test_inserts_total_and_blank_row_between_employees():
    rows = [
        _row('ALICE', run_act='2.00', run_est='4.00', run_adj='4.00', run_pct='200.0'),
        _row('BOB', run_act='1.00', run_est='1.00', run_adj='1.00', run_pct='100.0'),
    ]
    out = _employee_total_rows(rows, _report_def())
    # Alice's row, her Total row, a blank spacer, Bob's row, his Total row --
    # no trailing blank after the last employee.
    assert len(out) == 5
    assert out[0][0] == 'ALICE'
    assert out[1][0] == 'ALICE Total'
    assert out[2] == [''] * len(_COLUMNS)
    assert out[3][0] == 'BOB'
    assert out[4][0] == 'BOB Total'


def test_groups_by_employee_code_not_display_name():
    # Regression guard (CodeRabbit, PR #344): two different real employees
    # can share an identical formatted display name -- confirmed live, two
    # such pairs exist in this shop's own Employee table (not reproduced
    # here verbatim -- real names aren't committed to this repo, see
    # .claude/CLAUDE.md). Grouping by the 'employee' (display name) column
    # alone would silently merge their rows into one wrong combined Total
    # whenever that happens to land them adjacent in the query's own
    # ORDER BY, exactly the shape here.
    rows = [
        _row('SAMPLE OPERATOR', employee_code='SAMP01', run_act='2.00', run_est='4.00',
             run_adj='4.00', run_pct='200.0'),
        _row('SAMPLE OPERATOR', employee_code='SAMP02', run_act='1.00', run_est='1.00',
             run_adj='1.00', run_pct='100.0'),
    ]
    out = _employee_total_rows(rows, _report_def())
    col = {name: i for i, name in enumerate(_COLUMNS)}
    # Two separate Total rows, one per real employee -- not one combined
    # "SAMPLE OPERATOR Total" merging both codes' hours together.
    totals = [r for r in out if r[col['employee']] == 'SAMPLE OPERATOR Total']
    assert len(totals) == 2
    assert {t[col['run_act_hrs']] for t in totals} == {'2.00', '1.00'}


def test_total_row_sums_act_est_adj_and_recomputes_pct_as_ratio_of_sums():
    rows = [
        _row('ALICE', run_act='2.00', run_est='4.00', run_adj='4.00', run_pct='200.0'),
        _row('ALICE', run_act='3.00', run_est='3.00', run_adj='1.00', run_pct='33.3'),
    ]
    out = _employee_total_rows(rows, _report_def())
    total = out[-1]
    col = {name: i for i, name in enumerate(_COLUMNS)}
    assert total[col['run_act_hrs']] == '5.00'
    assert total[col['run_est_hrs']] == '7.00'
    assert total[col['run_adj_hrs']] == '5.00'
    # Ratio of sums (5/5*100=100), NOT an average of the two rows' own
    # % Eff values (which would be (200+33.3)/2 ~= 116.65).
    assert total[col['run_pct_eff']] == '100.0'


def test_unscoreable_rows_excluded_from_total_not_just_shown_blank():
    # A row with blank Adj Hrs (unscoreable, e.g. Est_Run_Hrs=0) must not
    # contribute its own Act/Est Hrs to the Total row -- same "exclude from
    # both sides of the ratio" rule already applied in Summary's aggregate.
    rows = [
        _row('ALICE', run_act='2.00', run_est='4.00', run_adj='4.00', run_pct='200.0'),
        _row('ALICE', run_act='999.00', run_est='0.00', run_adj='', run_pct=''),  # unscoreable
    ]
    out = _employee_total_rows(rows, _report_def())
    total = out[-1]
    col = {name: i for i, name in enumerate(_COLUMNS)}
    assert total[col['run_act_hrs']] == '2.00'  # not '1001.00'
    assert total[col['run_pct_eff']] == '200.0'


def test_blank_pct_eff_when_nothing_scoreable():
    rows = [_row('ALICE', run_act='5.00', run_est='0.00', run_adj='', run_pct='')]
    out = _employee_total_rows(rows, _report_def())
    total = out[-1]
    col = {name: i for i, name in enumerate(_COLUMNS)}
    assert total[col['run_act_hrs']] == ''
    assert total[col['run_pct_eff']] == ''


def test_pct_indirect_copied_from_first_row_not_resummed():
    rows = [_row('ALICE', pct_indirect='12.5'), _row('ALICE', pct_indirect='12.5')]
    out = _employee_total_rows(rows, _report_def())
    col = {name: i for i, name in enumerate(_COLUMNS)}
    assert out[-1][col['pct_indirect']] == '12.5'


def test_empty_input_returns_empty():
    assert _employee_total_rows([], _report_def()) == []


def test_noop_when_columns_are_not_detail_shaped():
    # A report without the expected Setup/Run columns (e.g. a plain job-list
    # report) must pass through unchanged rather than crash.
    simple_def = ReportDefinition(
        name="Simple", columns=('date', 'customer', 'job', 'description', 'status'),
        headers=['Date', 'Customer', 'Job', 'Description', 'Status'],
        fetch=lambda settings, filters: [], uses_employee_total_rows=True,
    )
    rows = [['2026-01-01', 'ACME', '123', 'Widget', 'Active']]
    assert _employee_total_rows(rows, simple_def) == rows


def test_write_xlsx_report_sheet_skips_excel_table_when_uses_employee_total_rows():
    # An Excel Table requires a uniform rectangular range -- the inserted
    # Total/blank rows break that, so a report with this flag on must not
    # get wrapped in one at all.
    wb = Workbook()
    ws = wb.active
    report_def = _report_def()
    rows = _employee_total_rows(
        [_row('ALICE', run_act='2.00', run_est='4.00', run_adj='4.00', run_pct='200.0')], report_def)
    _write_xlsx_report_sheet(
        ws, title_block=[['Title'], []], headers=list(_COLUMNS),
        data_rows=rows, report_def=report_def, table_name='ReportTable',
    )
    assert len(ws.tables) == 0


def test_write_xlsx_report_sheet_still_builds_table_when_flag_is_off():
    wb = Workbook()
    ws = wb.active
    report_def = _report_def(uses_employee_total_rows=False)
    rows = [_row('ALICE', run_act='2.00', run_est='4.00', run_adj='4.00', run_pct='200.0')]
    _write_xlsx_report_sheet(
        ws, title_block=[['Title'], []], headers=list(_COLUMNS),
        data_rows=rows, report_def=report_def, table_name='ReportTable',
    )
    assert 'ReportTable' in ws.tables
