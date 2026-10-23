"""
Reporting Module - JobBOSS DB reporting (experimental)

Pulls job data straight out of a JobBOSS SQL Server instead of showing
placeholder sample data. Connection details are configured in Settings >
Advanced Settings. Generic UI glue only -- every report type's query, row
shape, and filter requirements live in modules/reporting/reports/ (one file
per report type, see that package's __init__.py for the registry); this
file has no per-report-type branches, so adding a report type never
requires editing it.
"""

import sys
import csv
import html
import re
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Optional
from PyQt6.QtCore import QDate, QThread, pyqtSignal, QObject, QEvent, Qt, QUrl
from PyQt6.QtGui import QColor, QDesktopServices
from PyQt6.QtWidgets import (
    QWidget, QTableWidgetItem, QFileDialog, QHeaderView, QAbstractItemView,
    QDateEdit, QCalendarWidget, QDialog, QDialogButtonBox, QListWidget,
    QListWidgetItem, QPushButton, QVBoxLayout, QHBoxLayout, QComboBox, QFormLayout
)
from PyQt6 import uic

from core.base_module import BaseModule
from experimental.db_integration import is_configured, test_connection, fetch_customers, fetch_employees
from modules.reporting.reports import REPORTS, REPORT_TYPES

# Exported cells (Description/Notes/Customer/PO etc.) are free text sourced
# from JobBOSS -- neutralize values Excel/CSV readers treat as formula
# triggers (a note starting with "=" becomes a live formula on open; "+"/"-"/
# "@"/tab/CR are also sniffed by Excel's CSV importer). Prefixing with "'"
# is the standard mitigation (matches how Excel itself stores a manually
# typed "'=foo" as literal text).
_FORMULA_TRIGGER_CHARS = ('=', '+', '-', '@', '\t', '\r')


def _neutralize_formula_cell(value: str) -> str:
    if not value or not value.startswith(_FORMULA_TRIGGER_CHARS):
        return value
    # A plain negative number ("-5.2") is not a formula-injection risk --
    # only "-"-leading *text* is. Exempting it keeps numeric columns that
    # can go negative (e.g. a revenue credit) from being corrupted into
    # text cells; every other trigger char is still unconditionally
    # neutralized, since none of them are legitimate in a formatted number.
    if value.startswith('-'):
        try:
            float(value)
            return value
        except ValueError:
            pass
    return f"'{value}"


# Deliberately narrower than what int()/float() themselves accept: no
# exponent notation ("1e5"), no "nan"/"inf"/"infinity" text, no leading "+".
# float("nan")/float("inf") both succeed in plain Python, which would
# silently turn a literal text value spelled like one of those words into a
# special float instead of leaving it as text.
_SIMPLE_NUMBER_RE = re.compile(r'^-?\d+(\.\d+)?$')


def _numeric_or_text(value: str):
    """A numeric-looking cell value as a real int/float, or the original
    string if it isn't one.

    _export_table_rows() pulls everything from report_table as .text() --
    always a str, even "103.90" -- and openpyxl stores a plain str exactly
    as given, as a text cell, regardless of its content (unlike Excel's own
    CSV-open heuristic, which infers numbers from unquoted plain text).
    Every numeric export column (hours, quantities, % Eff) was landing in
    Excel as text: the "number stored as text" warning on every cell, and
    SUM() across the column silently returning 0 since there are no real
    number cells to add. int() tried first so whole numbers (e.g. Run Qty)
    don't round-trip display as "5.0". XLSX-only -- CSV has no native type
    system, so a plain "103.90" there is already correctly inferred as a
    number when Excel opens the file.

    A whole-number value only converts if it round-trips back to the exact
    same string: "30309" (a Job number) becomes int 30309 -- fine to sort/
    filter numerically, and str(30309) == "30309" so nothing is lost -- but
    "0012345" (a PO/part number with a real leading zero) does NOT round-
    trip (str(12345) == "12345" != "0012345") and stays text, since int()
    would otherwise silently discard a digit that's part of the real
    identifier, not a formatting artifact.
    """
    if not value or not _SIMPLE_NUMBER_RE.match(value):
        return value
    if '.' in value:
        return float(value)
    as_int = int(value)
    return as_int if str(as_int) == value else value


def _default_export_filename(report_type: Optional[str], fmt: str) -> str:
    """"JD_<ReportName>.<fmt>" -- e.g. "Job Report" -> "JD_JobReport.xlsx",
    "Employee Efficiency - Detail" -> "JD_EmployeeEfficiencyDetail.xlsx".
    Replaces the old timestamp-based "report_20261021_120144.xlsx" name,
    which told a user nothing about which report a file was without opening
    it -- especially awkward once several exports pile up in one folder.

    Each run of letters/digits in the report name is title-cased and
    concatenated (regex split on everything else: spaces, "-", etc.), so
    capitalization is normalized rather than copied verbatim -- "Jobs by
    Customer" becomes "JobsByCustomer", not "JobsbyCustomer". Falls back to
    a plain "Report" name if report_type is somehow unset (export_report()
    always sets it before this is called, but this stays a safe default
    rather than crashing on None).
    """
    words = re.findall(r'[A-Za-z0-9]+', report_type or '') or ['Report']
    return f"JD_{''.join(w.capitalize() for w in words)}.{fmt}"


def _group_rows_by_department(data_rows: list, department_col: int) -> list:
    """Group already-exported row lists by their Department column value,
    preserving first-seen order -- matches the SQL's own ORDER BY, so rows
    for one department stay in that relative order within their group.

    Backs the "By Department" export checkbox (base.ReportDefinition.
    uses_department_tabs): _export_xlsx() writes one worksheet tab per
    group this returns, instead of one combined sheet. A blank Department
    value (a work center with that column unset in JobBOSS) gets its own
    group rather than being silently dropped or merged into another one.
    """
    groups: dict = {}
    order = []
    for row in data_rows:
        department = row[department_col] or '(No Department)'
        if department not in groups:
            groups[department] = []
            order.append(department)
        groups[department].append(row)
    return [(department, groups[department]) for department in order]


_INVALID_SHEET_NAME_CHARS = frozenset(':\\/?*[]')


def _safe_sheet_name(name: str, used_names: set) -> str:
    """An Excel-legal worksheet name derived from a Department value.

    Excel worksheet names: <=31 chars, none of : \\ / ? * [ ], and unique
    within the workbook -- openpyxl raises if two sheets share a title,
    which truncation/sanitizing could cause even for two distinct
    department values. Mutates `used_names`, adding whatever name this
    call actually picks.
    """
    cleaned = ''.join('_' if c in _INVALID_SHEET_NAME_CHARS else c for c in (name or '')).strip()
    base = (cleaned or 'Sheet')[:31]
    candidate = base
    n = 2
    while candidate in used_names:
        suffix = f"_{n}"
        candidate = base[:31 - len(suffix)] + suffix
        n += 1
    used_names.add(candidate)
    return candidate


def _department_average_row(data_rows: list, average_columns: Optional[set] = None) -> list:
    """One extra row for a department's sheet: 'Department Average' in the
    first column, then the arithmetic mean of each other eligible column
    across every row in this group.

    average_columns restricts averaging to these column INDICES (the real
    metric columns -- hours, %, qty). None (the default, for callers with
    no column-type info) means every column is eligible. Restricting
    matters because an identifier column (job/operation/employee) can be
    all-numeric by coincidence (e.g. every value in a Job column is a bare
    number) and would otherwise get a meaningless average -- callers that
    know their column layout (_export_xlsx, via ReportDefinition.
    average_columns) should always pass the real metric-column set rather
    than relying on this default.

    Within an eligible column, a blank cell ('') means "not scoreable" (e.g.
    an Indirect operation's blank Adj Hrs) and is excluded from that
    column's average rather than voiding the whole column -- but any OTHER
    non-numeric value is a genuine anomaly and still blanks the whole
    column, same as before this function took an average_columns argument.
    """
    if not data_rows:
        return []
    width = len(data_rows[0])
    row = [''] * width
    row[0] = 'Department Average'
    eligible = (
        range(1, width) if average_columns is None
        else sorted(c for c in average_columns if 0 < c < width)
    )
    for col in eligible:
        values = []
        for r in data_rows:
            cell = r[col]
            if cell == '':
                continue
            try:
                values.append(float(cell))
            except (TypeError, ValueError):
                values = None
                break
        if values:
            row[col] = round(sum(values) / len(values), 2)
    return row


def _employee_total_rows(data_rows: list, report_def) -> list:
    """data_rows with an '<Employee> Total' row plus a blank spacer row
    inserted after each employee's consecutive block, matching JobBOSS's own
    native Detail report layout (one block per employee, each closed out by
    a total line). Relies on the rows already being grouped -- the report's
    own SQL ORDER BY puts one employee's rows together -- this does not
    re-sort or merge non-contiguous runs of the same employee.

    Groups by the 'employee_code' column when the report has one, not the
    displayed 'employee' name column -- two different real employees can
    share an identical formatted display name, which would otherwise merge
    their rows into one wrong combined Total.

    Within each employee's block, Setup/Run Act Hrs, Est Hrs, and Adj Hrs
    are each summed only across rows where that block's Adj Hrs is
    non-blank -- the same "not scoreable, exclude from both sides of the
    ratio" rule already applied per-cell (module docstring's "When JobBOSS
    blanks an Adj Hrs cell") and in Summary's own aggregate. % Eff is then
    the ratio of those sums, not an average of the rows' own % Eff values.
    % Indirect is a per-employee constant already identical on every one of
    that employee's rows, so it's copied from the first row rather than
    re-derived. A no-op (returns data_rows unchanged) if the column shape
    doesn't actually look like Detail's.
    """
    if not data_rows or 'employee' not in report_def.columns:
        return data_rows
    col = {name: i for i, name in enumerate(report_def.columns)}
    blocks = (
        ('setup_act_hrs', 'setup_est_hrs', 'setup_adj_hrs', 'setup_pct_eff'),
        ('run_act_hrs', 'run_est_hrs', 'run_adj_hrs', 'run_pct_eff'),
    )
    if not all(key in col for block in blocks for key in block):
        return data_rows
    name_idx = col['employee']
    # Group by the stable employee CODE when available, not the formatted
    # display name -- two different real employees can share an identical
    # "First Last" display name (confirmed live: two such pairs exist in
    # this shop's own Employee table), which would otherwise merge their
    # rows into one wrong combined Total whenever that happens to land them
    # adjacent in the query's own ORDER BY. (CodeRabbit, PR #344) Falls back
    # to the display name only for a report shape with no employee_code
    # column at all.
    group_idx = col.get('employee_code', name_idx)
    width = len(data_rows[0])

    def _num(value):
        try:
            return float(value)
        except (TypeError, ValueError):
            return None

    def _total_row(display_name: str, group: list) -> list:
        row = [''] * width
        row[name_idx] = f"{display_name} Total"
        for act_key, est_key, adj_key, pct_key in blocks:
            act_total = est_total = adj_total = 0.0
            scoreable = False
            for r in group:
                if r[col[adj_key]] == '':
                    continue
                scoreable = True
                act_total += _num(r[col[act_key]]) or 0.0
                est_total += _num(r[col[est_key]]) or 0.0
                adj_total += _num(r[col[adj_key]]) or 0.0
            if scoreable:
                # Strings, not raw floats -- every other cell in data_rows is
                # a string (same convention _numeric_or_text() downstream
                # assumes throughout this file), and this matches the same
                # f"{...:.2f}"/f"{...:.1f}" formatting _rows_from_detail()
                # already uses for these same columns.
                row[col[act_key]] = f"{act_total:.2f}"
                row[col[est_key]] = f"{est_total:.2f}"
                row[col[adj_key]] = f"{adj_total:.2f}"
                row[col[pct_key]] = f"{(adj_total / act_total * 100):.1f}" if act_total else ''
        if 'pct_indirect' in col:
            row[col['pct_indirect']] = group[0][col['pct_indirect']]
        return row

    result = []
    i, n = 0, len(data_rows)
    while i < n:
        group_key = data_rows[i][group_idx]
        j = i
        while j < n and data_rows[j][group_idx] == group_key:
            j += 1
        group = data_rows[i:j]
        result.extend(group)
        result.append(_total_row(group[0][name_idx], group))
        if j < n:
            result.append([''] * width)
        i = j
    return result


# A quoted, literal "%" -- NOT Excel's native unquoted "%" format code, which
# multiplies the stored value by 100 before displaying it. Our % Eff values
# are already stored in percent-point units (e.g. 136.9 meaning "136.9%"),
# not as a 0-1 fraction, so the native format would misdisplay 136.9 as
# "13690%". Rounds to the nearest whole percent for display only -- the full
# precision value is still what's stored in the cell and used by any
# formula/SUM() built on top of it in Excel. Matches JobBOSS's own real
# export exactly: its %Eff cells use the number format `#,##0[$%-409]`, a
# different-looking but equivalent "integer, literal % suffix, no auto-
# scaling" format (confirmed by inspecting a real export -- 85% of its %Eff
# cells store genuine non-integer values under that display mask).
_PCT_DISPLAY_FORMAT = '0"%"'


def _write_xlsx_report_sheet(ws, title_block: list, headers: list, data_rows: list, report_def,
                             average_row: Optional[list] = None, table_name: str = "ReportTable"):
    """Write one worksheet's worth of a report export: title block, header
    row, data rows, row_highlight() fill, bold styling, an Excel Table/
    AutoFilter, an optional trailing average row, and the report's footer
    lines -- everything _export_xlsx() used to do inline for its one-and-
    only sheet. Factored out so the "By Department" checkbox can call this
    once per department tab instead of once for the whole file; a plain,
    non-split export is just this called once.

    average_row (e.g. from _department_average_row()) is appended after the
    Table block, same placement reasoning as footer_lines below -- never
    part of the Table's own range, so Excel's AutoFilter/sort can't reorder
    it away from the bottom.
    """
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    for row in title_block:
        ws.append(row)
    header_row_idx = len(title_block) + 1

    # openpyxl's Table requires non-empty, unique headers -- some report
    # types repurpose blank columns for aggregate data (e.g. Top
    # Customers), so blanks get a placeholder name here; the on-screen
    # report_table is unaffected, this only changes what lands in the file.
    safe_headers = []
    seen = set()
    for i, h in enumerate(headers, start=1):
        name = h.strip() if h and h.strip() else f"Column{i}"
        if name in seen:
            name = f"{name}_{i}"
        seen.add(name)
        safe_headers.append(name)

    # Column KEYS ending in "_pct_eff" (e.g. setup_pct_eff, run_pct_eff) --
    # a naming convention shared by Detail and Summary -- get the display
    # format above. Data-driven by report_def.columns rather than a
    # hardcoded column letter so it applies to whichever report actually
    # has these columns, not just Employee Efficiency specifically.
    pct_eff_cols = (
        {i for i, key in enumerate(report_def.columns, start=1) if key.endswith('_pct_eff')}
        if report_def else set()
    )

    ws.append(safe_headers)
    for row in data_rows:
        ws.append([_numeric_or_text(v) for v in row])
        if pct_eff_cols:
            for col_idx in pct_eff_cols:
                ws.cell(row=ws.max_row, column=col_idx).number_format = _PCT_DISPLAY_FORMAT

    # Same row_highlight() flag as the on-screen report_table
    # (_on_report_rows) -- reapplied here from the exported string values
    # rather than read back off the QTableWidgetItem brushes, so this
    # stays correct even if someone exports without ever having the tab
    # open. Matched by column KEY (zip report_def.columns with the
    # exported row, same as _on_report_rows), not by header label -- one
    # shared way to find "which column" instead of two.
    if report_def and report_def.row_highlight and data_rows:
        red_fill = PatternFill(start_color='FFCCCC', end_color='FFCCCC', fill_type='solid')
        for i, row in enumerate(data_rows):
            row_dict = dict(zip(report_def.columns, row))
            highlight_key = report_def.row_highlight(row_dict)
            if highlight_key and highlight_key in report_def.columns:
                col_idx = report_def.columns.index(highlight_key)
                ws.cell(row=header_row_idx + 1 + i, column=col_idx + 1).fill = red_fill

    ws['A1'].font = Font(bold=True, size=12)
    for cell in ws[header_row_idx]:
        cell.font = Font(bold=True)
    for col_cells in ws.columns:
        width = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
        ws.column_dimensions[col_cells[0].column_letter].width = max(10, min(40, width + 2))

    # Wrap header + data in a real Excel Table with its AutoFilter, so the
    # exported file opens already sortable/filterable per column in Excel
    # -- not just a plain value dump. Skipped for a header-only (no data
    # rows) export: nothing to sort/filter, and Excel's Table requires at
    # least one row below the header. Also skipped when uses_employee_
    # total_rows inserted Total/blank rows into data_rows -- an Excel Table
    # range requires every row to be real, uniformly-shaped data, which the
    # inserted rows aren't.
    uses_total_rows = bool(report_def and report_def.uses_employee_total_rows)
    if len(data_rows) > 0 and not uses_total_rows:
        last_col_letter = get_column_letter(len(safe_headers))
        table_ref = f"A{header_row_idx}:{last_col_letter}{ws.max_row}"
        table = Table(displayName=table_name, ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)

    # Both appended after the Table block (table_ref above was already
    # fixed to header+data only) and after column-width sizing, so neither
    # a long average value nor a long footer line blows out every column's
    # width.
    if average_row is not None:
        # ws.max_row is computed from populated cells, not rows appended --
        # ws.append([]) (the blank spacer) writes no cells at all, so it
        # does NOT advance max_row even though it does occupy a real row.
        # Computing avg_row_idx as max_row + 1 right after that append (the
        # original code here) landed one row too early -- on the blank
        # spacer, not the average row -- so neither the bold font nor this
        # number format were ever actually reaching the real average row.
        # Reading max_row AFTER appending average_row itself sidesteps the
        # gap entirely, since that row always has real cell content (at
        # least the "Department Average" label).
        ws.append([])
        ws.append(average_row)
        avg_row_idx = ws.max_row
        for cell in ws[avg_row_idx]:
            cell.font = Font(bold=True)
        for col_idx in pct_eff_cols:
            ws.cell(row=avg_row_idx, column=col_idx).number_format = _PCT_DISPLAY_FORMAT

    if report_def and report_def.footer_lines:
        ws.append([])
        for line in report_def.footer_lines:
            ws.append([line])


# Fill color for whatever cell a report's row_highlight() flags (see
# modules/reporting/reports/base.py) -- currently only Job Report's
# Scheduled-End-past-Promise-Date flag, mirroring Report Fixer's Fix &
# Export red-highlight rule for the same comparison.
_LATE_FILL_COLOR = QColor(255, 204, 204)

# Cap on auto-fit column width (px) -- without it, one row with an unusually
# long Description/Notes value would stretch that single column wide enough
# to push every other column off-screen. The horizontal scrollbar (Interactive
# resize mode already enables it) covers genuinely-wide content beyond this;
# the user can still drag a capped column wider by hand if they want to see
# the full value without scrolling.
_MAX_AUTO_COLUMN_WIDTH = 300


class _CustomerPickerDialog(QDialog):
    """Pick zero, one, or several items for a report filter.

    Generic (value, label) checklist -- reused for both the Customer picker
    (value == label, a plain customer name straight off fetch_customers())
    and the Employee picker (value = the 6-char Employee code the SQL filter
    actually matches on, label = a human-readable "Name (Code)" display
    string off fetch_employees()); `title`/`noun` just relabel the dialog
    for whichever filter opened it.

    Zero checked means "All", not "none" -- Clear really clears back to All,
    it doesn't produce an empty/no-match report.
    """

    def __init__(self, items: list, selected: list, parent=None,
                 title: str = "Select Customers", noun: str = "Customers"):
        super().__init__(parent)
        self.setWindowTitle(title)
        self.resize(320, 420)
        layout = QVBoxLayout(self)

        self.list_widget = QListWidget()
        selected_set = set(selected)
        for value, label in items:
            item = QListWidgetItem(label)
            item.setData(Qt.ItemDataRole.UserRole, value)
            item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
            item.setCheckState(Qt.CheckState.Checked if value in selected_set else Qt.CheckState.Unchecked)
            self.list_widget.addItem(item)
        layout.addWidget(self.list_widget)

        btn_row = QHBoxLayout()
        select_all_btn = QPushButton("Select All")
        select_all_btn.clicked.connect(lambda: self._set_all(Qt.CheckState.Checked))
        clear_btn = QPushButton(f"Clear (All {noun})")
        clear_btn.clicked.connect(lambda: self._set_all(Qt.CheckState.Unchecked))
        btn_row.addWidget(select_all_btn)
        btn_row.addWidget(clear_btn)
        btn_row.addStretch()
        layout.addLayout(btn_row)

        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def _set_all(self, state: Qt.CheckState):
        for i in range(self.list_widget.count()):
            self.list_widget.item(i).setCheckState(state)

    def selected_values(self) -> list:
        return [
            self.list_widget.item(i).data(Qt.ItemDataRole.UserRole)
            for i in range(self.list_widget.count())
            if self.list_widget.item(i).checkState() == Qt.CheckState.Checked
        ]


class _ExportOptionsDialog(QDialog):
    """Format + Destination picker shown before every export, mirroring
    JobBOSS's own Crystal Reports export dialog (Format / Destination
    dropdowns) rather than letting the OS Save dialog's filter list be the
    only format choice. Defaults (Excel / Save to disk) reproduce this
    module's pre-existing export_report() behavior exactly, so accepting
    without changing anything is a no-op change for existing users.
    """

    _FORMATS = [("xlsx", "Excel (.xlsx)"), ("csv", "CSV (.csv)"), ("pdf", "PDF (.pdf)")]
    _DESTINATIONS = [("disk", "Save to disk"), ("app", "Open with default program")]

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Export Report")
        layout = QFormLayout(self)

        self.format_combo = QComboBox()
        for value, label in self._FORMATS:
            self.format_combo.addItem(label, value)
        layout.addRow("Format:", self.format_combo)

        self.destination_combo = QComboBox()
        for value, label in self._DESTINATIONS:
            self.destination_combo.addItem(label, value)
        layout.addRow("Destination:", self.destination_combo)

        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addRow(buttons)

    def selected_format(self) -> str:
        return self.format_combo.currentData()

    def selected_destination(self) -> str:
        return self.destination_combo.currentData()


class _DateCalendarPopupFilter(QObject):
    """Double-click a QDateEdit to drop down a small calendar for picking a
    date, in addition to Qt's built-in dropdown-arrow click (calendarPopup).
    Built on public QCalendarWidget/QWidget APIs rather than QDateEdit's
    private popup internals, which aren't safe to trigger programmatically.

    QAbstractSpinBox (QDateEdit's base) routes mouse events for the visible
    text to its internal child QLineEdit, not to the QDateEdit itself -- so
    this must be installed on date_edit.lineEdit(), not date_edit directly,
    or a real double-click never reaches eventFilter() at all.
    """

    def __init__(self, date_edit: QDateEdit, parent=None):
        super().__init__(parent)
        self._date_edit = date_edit

    def eventFilter(self, obj, event) -> bool:
        if event.type() == QEvent.Type.MouseButtonDblClick:
            self._show_popup()
            return True
        return False

    def _show_popup(self):
        date_edit = self._date_edit
        calendar = QCalendarWidget()
        calendar.setWindowFlags(Qt.WindowType.Popup)
        calendar.setSelectedDate(date_edit.date())

        def pick(qdate):
            date_edit.setDate(qdate)
            calendar.close()

        calendar.clicked.connect(pick)
        calendar.activated.connect(pick)
        calendar.move(date_edit.mapToGlobal(date_edit.rect().bottomLeft()))
        calendar.show()
        self._popup = calendar  # keep alive while shown; see CODING_NOTES.md on non-modal dialogs


class _ConnectWorker(QThread):
    """One-shot connection test, off the GUI thread (mirrors main.py's _PluginInstallWorker)."""

    finished_ok = pyqtSignal(bool, str)

    def __init__(self, settings: dict):
        super().__init__()
        self._settings = settings

    def run(self):
        ok, message = test_connection(self._settings)
        self.finished_ok.emit(ok, message)


class _CustomerListWorker(QThread):
    """Fetches the distinct customer list to populate the Customer filter combo."""

    success = pyqtSignal(list)
    error = pyqtSignal(str)

    def __init__(self, settings: dict):
        super().__init__()
        self._settings = settings

    def run(self):
        try:
            customers = fetch_customers(self._settings)
        except Exception as exc:
            self.error.emit(f"Customer list query failed ({type(exc).__name__})")
        else:
            self.success.emit(customers)


class _EmployeeListWorker(QThread):
    """Fetches the active employee list to populate the Employee filter picker."""

    success = pyqtSignal(list)
    error = pyqtSignal(str)

    def __init__(self, settings: dict):
        super().__init__()
        self._settings = settings

    def run(self):
        try:
            employees = fetch_employees(self._settings)
        except Exception as exc:
            self.error.emit(f"Employee list query failed ({type(exc).__name__})")
        else:
            self.success.emit(employees)


class _ReportWorker(QThread):
    """Runs one report query off the GUI thread."""

    success = pyqtSignal(list)
    error = pyqtSignal(str)

    def __init__(self, settings: dict, report_type: str, filters: dict):
        super().__init__()
        self._settings = settings
        self._report_type = report_type
        self._filters = filters

    def run(self):
        try:
            if not is_configured(self._settings):
                raise RuntimeError("JobBOSS DB is not configured")
            rows = REPORTS[self._report_type].fetch(self._settings, self._filters)
        except Exception as exc:
            # Only the exception type, not str(exc) -- same rule as
            # experimental/db_integration.py's test_connection().
            self.error.emit(f"Report query failed ({type(exc).__name__})")
        else:
            self.success.emit(rows)


class ReportingModule(BaseModule):
    """Module for generating and exporting JobBOSS reports"""

    def __init__(self):
        super().__init__()
        self._widget = None
        # Widget references
        self.report_type_combo = None
        self.report_table = None
        self.report_status_label = None
        self.generate_report_btn = None
        self.db_status_indicator = None
        self.customer_filter_btn = None
        self.employee_filter_btn = None
        self.department_split_check = None
        self.exclude_assemblies_check = None
        self.report_start_date = None
        self.report_end_date = None
        self._available_customers = []
        self._selected_customers = []
        self._available_employees = []
        self._selected_employees = []
        self._last_department_split = False

        self._db_ready = False
        self._connect_worker = None
        self._report_worker = None
        self._customer_worker = None
        self._employee_worker = None

        # Metadata for whatever report is currently in report_table, captured
        # at generate time so Export reflects what was actually run even if
        # the combo/filters have changed since (not whatever they show now).
        self._last_report_type = None
        self._last_filters = {}
        self._last_generated_at = None

    def get_name(self) -> str:
        return "Reports (Beta)"

    def get_order(self) -> int:
        return 80  # Eighth tab

    def is_experimental(self) -> bool:
        return True  # This module is experimental

    def initialize(self, app_context):
        super().initialize(app_context)

    def get_widget(self) -> QWidget:
        if self._widget is None:
            self._widget = self._create_widget()
        return self._widget

    def _create_widget(self) -> QWidget:
        """Create the reporting tab widget"""
        widget = QWidget()

        # Load UI file
        ui_file = self._get_ui_path('reporting/ui/reporting_tab.ui')
        uic.loadUi(ui_file, widget)

        # Store widget references
        self.report_type_combo = widget.report_type_combo
        self.report_table = widget.report_table
        self.report_status_label = widget.report_status_label
        self.generate_report_btn = widget.generate_report_btn
        self.db_status_indicator = widget.db_status_indicator
        self.customer_filter_btn = widget.customer_filter_btn
        self.employee_filter_btn = widget.employee_filter_btn
        self.department_split_check = widget.department_split_check
        self.exclude_assemblies_check = widget.exclude_assemblies_check
        self.report_start_date = widget.report_start_date
        self.report_end_date = widget.report_end_date
        # End defaults to today, not "(No Filter)" -- "through today" is the
        # overwhelmingly common upper bound, so this saves a click on every
        # date-filtered report. Start has no equally obvious default, so it
        # stays at reporting_tab.ui's "(No Filter)" sentinel (minimumDate)
        # until the user actually picks one.
        self.report_end_date.setDate(QDate.currentDate())

        self._start_date_popup_filter = _DateCalendarPopupFilter(self.report_start_date, widget)
        self.report_start_date.lineEdit().installEventFilter(self._start_date_popup_filter)
        self._end_date_popup_filter = _DateCalendarPopupFilter(self.report_end_date, widget)
        self.report_end_date.lineEdit().installEventFilter(self._end_date_popup_filter)

        # Setup table properties
        self.report_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.report_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)

        # Connect signals
        self.db_status_indicator.clicked.connect(self.connect_to_db)
        self.generate_report_btn.clicked.connect(self.generate_report)
        self.customer_filter_btn.clicked.connect(self._open_customer_picker)
        self.employee_filter_btn.clicked.connect(self._open_employee_picker)
        widget.export_report_btn.clicked.connect(self.export_report)

        # Populated from the reports registry, not reporting_tab.ui's XML --
        # adding a report type means adding a file to modules/reporting/reports/,
        # never touching this .ui file.
        self.report_type_combo.clear()
        self.report_type_combo.addItems(REPORT_TYPES)

        self.report_type_combo.currentTextChanged.connect(self._on_report_type_changed)
        self._on_report_type_changed(self.report_type_combo.currentText())

        # Auto-connect as soon as the tab is opened -- no manual "Connect"
        # step anymore. get_widget() caches this QWidget, so _create_widget()
        # (and this call) only run once per app session; db_status_indicator
        # is clickable to retry if this attempt fails or the server was
        # briefly unreachable.
        self.connect_to_db()

        return widget

    def _on_report_type_changed(self, report_type: str):
        self._update_filter_enabled_state(report_type)
        # Only reset rows/headers immediately if nothing's in flight -- if a
        # report is still running, leave the old (still internally consistent)
        # headers/rows in place rather than resetting to the new type's empty
        # layout; _on_report_rows() applies the correct headers for whichever
        # report actually completes, using the type that was really queried.
        if self._report_worker is None or not self._report_worker.isRunning():
            self.report_table.setRowCount(0)
            self._update_table_headers(report_type)

    def _update_filter_enabled_state(self, report_type: str):
        """Grey out the Customer/Employee/Date filters when the selected
        report type doesn't use them, instead of silently ignoring whatever's
        set.
        """
        report_def = REPORTS[report_type]
        self.customer_filter_btn.setEnabled(report_def.uses_customer_filter)
        self.employee_filter_btn.setEnabled(report_def.uses_employee_filter)
        self.department_split_check.setEnabled(report_def.uses_department_tabs)
        self.exclude_assemblies_check.setEnabled(report_def.uses_exclude_assemblies)
        self.report_start_date.setEnabled(report_def.uses_dates)
        self.report_end_date.setEnabled(report_def.uses_dates)

    def _report_columns(self, report_type: str) -> tuple:
        """The row-dict keys to pull into report_table, in column order, for report_type."""
        return REPORTS[report_type].columns

    def _update_table_headers(self, report_type: str):
        """Apply the headers/column count/resize mode the registered report
        declares for itself -- every report type owns its own headers now, no
        shared fallback.
        """
        report_def = REPORTS[report_type]
        headers = report_def.headers

        self.report_table.setColumnCount(len(headers))
        for col, label in enumerate(headers):
            self.report_table.setHorizontalHeaderItem(col, QTableWidgetItem(label))

        # Stretch-to-fit reads fine for 5 columns but crams Job Report's 16
        # into illegibly narrow slivers -- Interactive lets the user resize
        # and scroll instead.
        resize_mode = (QHeaderView.ResizeMode.Interactive if report_def.resize_mode == "interactive"
                       else QHeaderView.ResizeMode.Stretch)
        self.report_table.horizontalHeader().setSectionResizeMode(resize_mode)

    def _get_ui_path(self, relative_path: str) -> Path:
        """Get path to UI file"""
        if getattr(sys, 'frozen', False):
            application_path = Path(sys._MEIPASS)
        else:
            application_path = Path(__file__).parent.parent.parent

        ui_file = application_path / 'modules' / relative_path
        if not ui_file.exists():
            raise FileNotFoundError(f"UI file not found: {ui_file}")
        return ui_file

    def _jobboss_settings(self) -> dict:
        """Collect the non-secret JobBOSS DB connection settings; the password
        itself is looked up from the OS credential store by db_integration.
        """
        return {
            'jobboss_db_host': self.app_context.get_setting('jobboss_db_host', ''),
            'jobboss_db_port': self.app_context.get_setting('jobboss_db_port', ''),
            'jobboss_db_name': self.app_context.get_setting('jobboss_db_name', ''),
            'jobboss_db_user': self.app_context.get_setting('jobboss_db_user', ''),
        }

    # ==================== Database Connection ====================

    def connect_to_db(self):
        """Test the JobBOSS DB connection and refresh the status indicator.

        Called once automatically when the tab first opens, and again on
        every click of db_status_indicator itself (the only way to retry,
        since get_widget() caches this QWidget -- _create_widget() and this
        initial call only run once per app session). No popup here even
        when unconfigured/unreachable -- that would fire on every tab open
        for anyone who hasn't set up DB access yet. generate_report() does
        its own is_configured() check and shows a real error dialog instead,
        since that's a deliberate user action, not something that just ran
        silently in the background.
        """
        # Guard against the auto-connect-on-open call racing a quick manual
        # click of db_status_indicator (now the only way to retry): replacing
        # self._connect_worker while the old one is still running drops its
        # only reference (it has no Qt parent), and Qt aborts the process on
        # "destroyed while thread is still running" once that worker's run()
        # eventually finishes. Same pattern as _load_customer_list() below.
        if self._connect_worker is not None and self._connect_worker.isRunning():
            return

        settings = self._jobboss_settings()
        if not is_configured(settings):
            self._db_ready = False
            self._set_db_status_indicator(
                "#999", "JobBOSS database not configured -- see Settings > Advanced Settings. Click to retry."
            )
            return

        self._set_db_status_indicator("#999", "Connecting to JobBOSS database...")

        self._connect_worker = _ConnectWorker(settings)
        self._connect_worker.finished_ok.connect(self._on_connect_result)
        self._connect_worker.start()

    def _on_connect_result(self, ok: bool, message: str):
        self._db_ready = ok
        if ok:
            self._set_db_status_indicator("#2e8b2e", "Connected to JobBOSS database. Click to re-check.")
            self.log_message("Connected to JobBOSS DB")
            self._load_customer_list()
            self._load_employee_list()
        else:
            self._set_db_status_indicator("#c0392b", f"{message} Click to retry.")

    def _set_db_status_indicator(self, color: str, tooltip: str):
        self.db_status_indicator.setStyleSheet(
            f"QPushButton {{ background-color: {color}; border-radius: 7px; border: none; }}"
        )
        self.db_status_indicator.setToolTip(tooltip)

    def _load_customer_list(self):
        """Populate the Customer filter combo from JobBOSS after a successful connect."""
        # Guard against Connect -> Disconnect -> Connect faster than the
        # distinct-customer query finishes: replacing self._customer_worker
        # while the old one is still running drops its only reference (it has
        # no Qt parent), and Qt aborts the process on "destroyed while thread
        # is still running" once that worker's run() eventually finishes.
        if self._customer_worker is not None and self._customer_worker.isRunning():
            return
        self._customer_worker = _CustomerListWorker(self._jobboss_settings())
        self._customer_worker.success.connect(self._on_customers_loaded)
        self._customer_worker.error.connect(
            lambda msg: self.log_message(f"Reporting: {msg}")
        )
        self._customer_worker.start()

    def _on_customers_loaded(self, customers: list):
        self._available_customers = customers
        # Drop any previously-selected customer that's no longer in the list
        # (e.g. reconnected to a different JobBOSS DB) rather than silently
        # filtering on a name that can never match again.
        customer_set = set(customers)
        self._selected_customers = [c for c in self._selected_customers if c in customer_set]
        self._update_customer_filter_button_text()

    def _open_customer_picker(self):
        items = [(name, name) for name in self._available_customers]
        dialog = _CustomerPickerDialog(
            items, self._selected_customers, self._widget,
            title="Select Customers", noun="Customers")
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self._selected_customers = dialog.selected_values()
            self._update_customer_filter_button_text()

    def _update_customer_filter_button_text(self):
        n = len(self._selected_customers)
        if n == 0:
            text = "Customers: All"
        elif n == 1:
            text = f"Customer: {self._selected_customers[0]}"
        else:
            text = f"Customers: {n} selected"
        self.customer_filter_btn.setText(text)

    def _load_employee_list(self):
        """Populate the Employee filter picker from JobBOSS after a successful connect."""
        if self._employee_worker is not None and self._employee_worker.isRunning():
            return
        self._employee_worker = _EmployeeListWorker(self._jobboss_settings())
        self._employee_worker.success.connect(self._on_employees_loaded)
        self._employee_worker.error.connect(
            lambda msg: self.log_message(f"Reporting: {msg}")
        )
        self._employee_worker.start()

    def _on_employees_loaded(self, employees: list):
        self._available_employees = employees
        codes = {e['code'] for e in employees}
        self._selected_employees = [c for c in self._selected_employees if c in codes]
        self._update_employee_filter_button_text()

    def _open_employee_picker(self):
        items = [(e['code'], f"{e['name']} ({e['code']})") for e in self._available_employees]
        dialog = _CustomerPickerDialog(
            items, self._selected_employees, self._widget,
            title="Select Employees", noun="Employees")
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self._selected_employees = dialog.selected_values()
            self._update_employee_filter_button_text()

    def _employee_name_by_code(self) -> dict:
        return {e['code']: e['name'] for e in self._available_employees}

    def _update_employee_filter_button_text(self):
        n = len(self._selected_employees)
        if n == 0:
            text = "Employees: All"
        elif n == 1:
            name = self._employee_name_by_code().get(self._selected_employees[0], self._selected_employees[0])
            text = f"Employee: {name}"
        else:
            text = f"Employees: {n} selected"
        self.employee_filter_btn.setText(text)

    @staticmethod
    def _picked_date(date_edit: QDateEdit):
        """The QDate as a Python date, or None if still at the "(No Filter)"
        sentinel (date_edit's minimumDate, set in reporting_tab.ui).
        """
        if date_edit.date() == date_edit.minimumDate():
            return None
        return date_edit.date().toPyDate()

    # ==================== Report Generation ====================

    def generate_report(self):
        """Generate a report by querying the JobBOSS DB on a background thread.

        One generic filter-building block driven by the report's declared
        flags (base.ReportDefinition) -- a new report type with any
        combination of these flags needs no new code here.
        """
        # A cheap, synchronous local-settings check, not the _db_ready flag --
        # that depends on the auto-connect test's async round-trip, which
        # may still be running or may have failed transiently even though
        # the DB is reachable now. Letting the real query attempt (below)
        # be the actual source of truth for reachability means a slow/flaky
        # first connect doesn't permanently block Generate Report; only a
        # genuinely unconfigured DB does, which deserves this clear message
        # instead of _on_report_error()'s generic "(ExceptionType)" text.
        if not is_configured(self._jobboss_settings()):
            self.show_error(
                "Not Configured",
                "Set the JobBOSS DB connection (Host, Database, Username, Password) "
                "in Settings > Advanced Settings first."
            )
            return

        report_type = self.report_type_combo.currentText()
        report_def = REPORTS[report_type]
        filters = {}

        if report_def.uses_customer_filter and self._selected_customers:
            filters['customers'] = list(self._selected_customers)
        if report_def.uses_employee_filter and self._selected_employees:
            filters['employees'] = list(self._selected_employees)
        if report_def.uses_exclude_assemblies and self.exclude_assemblies_check.isChecked():
            filters['exclude_assemblies'] = True

        if report_def.uses_dates:
            start = self._picked_date(self.report_start_date)
            end = self._picked_date(self.report_end_date)
            if report_def.dates_required:
                if start is None or end is None:
                    self.show_error(
                        "Invalid Date",
                        "Pick both a Start Date and End Date (double-click a field, or use its "
                        "dropdown arrow, to open the calendar)."
                    )
                    return
                filters['start_date'] = start
                filters['end_date'] = end
            else:
                if (start is None) != (end is None):
                    self.show_error(
                        "Invalid Date",
                        "Pick both Start Date and End Date, or leave both as \"(No Filter)\" for all time."
                    )
                    return
                if start is not None and end is not None:
                    filters['start_date'] = start
                    filters['end_date'] = end

        self.generate_report_btn.setEnabled(False)
        self.report_status_label.setText(f"Running '{report_type}'...")

        self._last_report_type = report_type
        self._last_filters = filters
        # Captured at generate time, not read fresh at export time -- the
        # checkbox can change (or the report type can switch, disabling it)
        # while this run's results still sit in report_table unexported.
        self._last_department_split = report_def.uses_department_tabs and self.department_split_check.isChecked()

        self._report_worker = _ReportWorker(self._jobboss_settings(), report_type, filters)
        self._report_worker.success.connect(self._on_report_rows)
        self._report_worker.error.connect(self._on_report_error)
        self._report_worker.start()

    def _on_report_rows(self, rows: list):
        self.generate_report_btn.setEnabled(True)

        # _last_report_type/_last_filters, not the live combo/field selection
        # -- both can change while this report is still running in the
        # background, and these rows were produced from whatever was actually
        # sent to the worker at generate_report() time.
        report_type = self._last_report_type
        report_def = REPORTS[report_type]
        columns = report_def.columns
        # Apply headers for the report type that was actually queried, not
        # whatever the combo currently shows -- otherwise a combo switch
        # while this query was running leaves stale, mismatched headers/
        # column-count in place and setItem()/the highlight lookup below can
        # silently truncate or index into a None cell.
        self._update_table_headers(report_type)

        self.report_table.setRowCount(0)
        for row in rows:
            r = self.report_table.rowCount()
            self.report_table.insertRow(r)
            for col, key in enumerate(columns):
                self.report_table.setItem(r, col, QTableWidgetItem(row.get(key, '')))
            highlight_key = report_def.row_highlight(row) if report_def.row_highlight else None
            if highlight_key:
                self.report_table.item(r, columns.index(highlight_key)).setBackground(_LATE_FILL_COLOR)

        # Auto-fit columns to their actual content (header + widest cell) now
        # that the table is populated -- Interactive resize mode's default
        # column widths don't reflect content at all, which is what made Job
        # Report's 16 columns read as compacted/truncated. No-op in Stretch
        # mode (the 5-column reports): Stretch always re-fills the viewport
        # width regardless, so this has no lasting visible effect there.
        self.report_table.resizeColumnsToContents()
        for col in range(self.report_table.columnCount()):
            if self.report_table.columnWidth(col) > _MAX_AUTO_COLUMN_WIDTH:
                self.report_table.setColumnWidth(col, _MAX_AUTO_COLUMN_WIDTH)

        self._last_generated_at = datetime.now()
        self.report_status_label.setText(f"Showing {len(rows)} record(s) for '{report_type}'")
        self.log_message(f"Generated report: {report_type} ({len(rows)} rows)")

    def _on_report_error(self, message: str):
        self.generate_report_btn.setEnabled(True)
        self.report_status_label.setText("Report failed.")
        self.show_error("Report Failed", message)

    def _date_range_label(self) -> str:
        start = self._last_filters.get('start_date')
        end = self._last_filters.get('end_date')
        if start and end:
            return f"{start.strftime('%Y-%m-%d')} to {end.strftime('%Y-%m-%d')}"
        report_def = REPORTS.get(self._last_report_type)
        if report_def and report_def.date_label_when_unset:
            return report_def.date_label_when_unset
        return "N/A"

    def _export_title_block(self) -> list:
        generated = self._last_generated_at
        report_def = REPORTS.get(self._last_report_type)
        rows = [
            ["Report:", self._last_report_type or ""],
            ["Date Range:", self._date_range_label()],
        ]
        if report_def and report_def.uses_customer_filter:
            customers = self._last_filters.get('customers')
            value = ", ".join(customers) if customers else "All"
            rows.append(["Customers:", _neutralize_formula_cell(value)])
        if report_def and report_def.uses_employee_filter:
            codes = self._last_filters.get('employees')
            name_by_code = self._employee_name_by_code()
            labels = [name_by_code.get(c, c) for c in codes] if codes else []
            value = ", ".join(labels) if labels else "All"
            rows.append(["Employees:", _neutralize_formula_cell(value)])
        rows.append(["Generated:", generated.strftime('%Y-%m-%d %H:%M:%S') if generated else ""])
        rows.append([])
        return rows

    def _export_table_rows(self) -> tuple:
        """(headers, data_rows) from report_table, as plain strings."""
        headers = [
            self.report_table.horizontalHeaderItem(col).text()
            for col in range(self.report_table.columnCount())
        ]
        data_rows = []
        for row in range(self.report_table.rowCount()):
            data_rows.append([
                self.report_table.item(row, col).text() if self.report_table.item(row, col) else ""
                for col in range(self.report_table.columnCount())
            ])
        return headers, data_rows

    _EXPORT_FILTERS = {
        'xlsx': "Excel Files (*.xlsx)",
        'csv': "CSV Files (*.csv)",
        'pdf': "PDF Files (*.pdf)",
    }
    _EXPORT_METHODS_BY_FORMAT = {
        'xlsx': '_export_xlsx',
        'csv': '_export_csv',
        'pdf': '_export_pdf',
    }

    def export_report(self):
        """Export report to Excel (.xlsx), CSV, or PDF -- to disk, or opened
        directly in the OS's default program for that file type.
        """
        if self.report_table.rowCount() == 0:
            self.show_error("No Data", "Generate a report first before exporting")
            return

        dialog = _ExportOptionsDialog(self._widget)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return
        fmt = dialog.selected_format()
        destination = dialog.selected_destination()
        default_name = _default_export_filename(self._last_report_type, fmt)

        if destination == 'disk':
            file_path, _selected_filter = QFileDialog.getSaveFileName(
                self._widget, "Export Report", default_name, self._EXPORT_FILTERS[fmt]
            )
            if not file_path:
                return
            if not file_path.lower().endswith(f'.{fmt}'):
                file_path += f'.{fmt}'
        else:
            # 'app' destination: no save-location prompt, write to a temp
            # file and hand it straight to the OS's default handler --
            # matches how JobBOSS's own "Application" export destination
            # behaves (opens directly, doesn't ask where to save first).
            # Needs a timestamp suffix here (unlike the Save dialog's
            # default_name above) -- re-exporting the same report type
            # twice in a row would otherwise reuse the exact same temp
            # path, and overwriting it while the first export is still open
            # in Excel fails with a permission error instead of opening a
            # fresh copy.
            stem, _, suffix = default_name.rpartition('.')
            temp_name = f"{stem}_{datetime.now().strftime('%H%M%S')}.{suffix}"
            file_path = str(Path(tempfile.gettempdir()) / temp_name)

        try:
            getattr(self, self._EXPORT_METHODS_BY_FORMAT[fmt])(file_path)
        except Exception as e:
            self.show_error("Export Failed", f"Failed to export report:\n{str(e)}")
            return

        if destination == 'app':
            if not QDesktopServices.openUrl(QUrl.fromLocalFile(file_path)):
                self.show_error(
                    "Couldn't Open File",
                    f"The report was exported, but no default program could be found to open it:\n{file_path}",
                )
            self.log_message(f"Exported report and opened: {file_path}")
        else:
            self.show_info("Export Successful", f"Report exported to:\n{file_path}")
            self.log_message(f"Exported report to: {file_path}")

    def _export_csv(self, file_path: str):
        headers, data_rows = self._export_table_rows()
        safe_rows = [[_neutralize_formula_cell(v) for v in row] for row in data_rows]
        report_def = REPORTS.get(self._last_report_type)
        # utf-8-sig, not the platform-default encoding (cp1252 on Windows) --
        # a customer name or note containing a character outside cp1252 would
        # otherwise raise UnicodeEncodeError and fail the export; the BOM also
        # makes Excel itself open the file as UTF-8 instead of guessing wrong.
        with open(file_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerows(self._export_title_block())
            writer.writerow(headers)
            writer.writerows(safe_rows)
            if report_def and report_def.footer_lines:
                writer.writerow([])
                writer.writerows([[line] for line in report_def.footer_lines])

    def _export_pdf(self, file_path: str):
        """Render the title block, table, and footer as HTML and print it
        to a PDF via Qt's own QTextDocument/QPrinter -- no extra dependency
        beyond the PyQt6 install this app already requires. No formula-
        injection neutralizing needed here (unlike CSV/xlsx): a PDF has no
        spreadsheet engine to execute a cell value as a formula, so only
        HTML-escaping (for the markup itself, e.g. a note containing "<"/"&")
        applies.
        """
        from PyQt6.QtGui import QPageLayout, QTextDocument
        from PyQt6.QtPrintSupport import QPrinter

        def esc(value) -> str:
            return html.escape(str(value)) if value not in (None, "") else ""

        title_block = self._export_title_block()
        headers, data_rows = self._export_table_rows()
        report_def = REPORTS.get(self._last_report_type)

        highlighted_cells = set()
        if report_def and report_def.row_highlight:
            for row_idx, row in enumerate(data_rows):
                row_dict = dict(zip(report_def.columns, row))
                highlight_key = report_def.row_highlight(row_dict)
                if highlight_key and highlight_key in report_def.columns:
                    highlighted_cells.add((row_idx, report_def.columns.index(highlight_key)))

        parts = ['<html><body style="font-family: sans-serif; font-size: 9pt;">']
        for i, row in enumerate(title_block):
            if not row:
                continue
            text = " ".join(esc(cell) for cell in row if cell not in (None, ""))
            style = "font-weight:bold; font-size:13pt;" if i == 0 else "margin:1px 0;"
            parts.append(f'<p style="{style}">{text}</p>')

        parts.append('<table border="1" cellspacing="0" cellpadding="4" style="border-collapse:collapse; width:100%;">')
        parts.append('<tr>' + ''.join(f'<th style="background:#dddddd;">{esc(h)}</th>' for h in headers) + '</tr>')
        for row_idx, row in enumerate(data_rows):
            cells = []
            for col_idx, value in enumerate(row):
                style = ' style="background:#ffcccc;"' if (row_idx, col_idx) in highlighted_cells else ''
                cells.append(f'<td{style}>{esc(value)}</td>')
            parts.append('<tr>' + ''.join(cells) + '</tr>')
        parts.append('</table>')

        if report_def and report_def.footer_lines:
            parts.append('<p>&nbsp;</p>')
            for line in report_def.footer_lines:
                parts.append(f'<p style="font-size:7.5pt; color:#444444;">{esc(line)}</p>')
        parts.append('</body></html>')

        document = QTextDocument()
        document.setHtml(''.join(parts))

        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(file_path)
        printer.setPageOrientation(QPageLayout.Orientation.Landscape)
        document.print(printer)

    def _export_xlsx(self, file_path: str):
        from openpyxl import Workbook

        title_block = self._export_title_block()
        headers, data_rows = self._export_table_rows()
        # Neutralizing here (not just in _export_csv) matters for xlsx too --
        # openpyxl auto-detects data_type='f' (formula) for any plain string
        # cell value starting with "=", so an unprefixed JobBOSS note would
        # become a *real* live formula, not just something Excel's CSV import
        # heuristically treats as one. Safe to do once, upfront: date/numeric
        # columns never start with a trigger char, so the highlight lookup
        # below (which reads these same data_rows by index) is unaffected.
        data_rows = [[_neutralize_formula_cell(v) for v in row] for row in data_rows]
        report_def = REPORTS.get(self._last_report_type)

        department_col = None
        if (self._last_department_split and report_def is not None
                and report_def.uses_department_tabs and 'department' in report_def.columns):
            department_col = report_def.columns.index('department')

        average_columns = None
        if report_def is not None and report_def.average_columns:
            average_columns = {
                report_def.columns.index(key)
                for key in report_def.average_columns if key in report_def.columns
            }

        uses_total_rows = bool(report_def and report_def.uses_employee_total_rows)

        wb = Workbook()
        if department_col is not None:
            used_sheet_names: set = set()
            for i, (department, rows) in enumerate(_group_rows_by_department(data_rows, department_col)):
                ws = wb.active if i == 0 else wb.create_sheet()
                ws.title = _safe_sheet_name(department, used_sheet_names)
                table_name = "ReportTable" if i == 0 else f"ReportTable{i + 1}"
                # Average computed from the clean per-operation rows, before
                # Total/blank rows are interleaved -- those rows would
                # otherwise double-count already-summed values into the
                # department's own average.
                avg_row = _department_average_row(rows, average_columns)
                sheet_rows = _employee_total_rows(rows, report_def) if uses_total_rows else rows
                _write_xlsx_report_sheet(
                    ws, title_block, headers, sheet_rows, report_def,
                    average_row=avg_row, table_name=table_name)
        else:
            ws = wb.active
            ws.title = "Report"
            if uses_total_rows:
                data_rows = _employee_total_rows(data_rows, report_def)
            _write_xlsx_report_sheet(ws, title_block, headers, data_rows, report_def)

        wb.save(file_path)

    def cleanup(self):
        """Cleanup resources.

        No persistent JobBOSS connection to close here -- experimental.
        db_integration.connect()/test_connection() each open a pytds
        connection in their own `with` block and close it before returning,
        every single call, so there's never a live session held open
        between calls for this to disconnect. What actually needs cleanup
        is the background workers: wait for any still in flight (most
        importantly _connect_worker, since a module teardown mid-connect
        could otherwise leave a QThread running after its receiver is gone)
        and reset local state so nothing stale (a green "connected"
        indicator with no backing worker) survives this module being torn
        down and rebuilt.
        """
        workers = (self._connect_worker, self._report_worker, self._customer_worker, self._employee_worker)
        for worker in workers:
            if worker is not None and worker.isRunning():
                worker.wait(2000)
        self._db_ready = False
